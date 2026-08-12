# -*- coding: utf-8 -*-
"""Dependency-isolated text extraction for OpenXnet office documents."""

from __future__ import annotations

import json
from pathlib import Path, PurePosixPath
import posixpath
import re
import sys
from typing import Callable
import zipfile

from py.worker_capability_contracts import (
    PORTABLE_DOCUMENT_FORMATS,
    SUPPORTED_DOCUMENT_FORMATS,
    WINDOWS_DOCUMENT_FORMATS,
)

DEFAULT_MAX_OUTPUT_BYTES = 32 * 1024 * 1024
DEFAULT_MAX_EPUB_ENTRIES = 10_000
DEFAULT_MAX_EPUB_UNCOMPRESSED_BYTES = 512 * 1024 * 1024


class DocumentEngine:
    """Extract plain text from supported document formats with bounded output."""

    def __init__(self, *, max_output_bytes: int = DEFAULT_MAX_OUTPUT_BYTES) -> None:
        """Create a stateless parser with one UTF-8 output-size limit."""

        self._max_output_bytes = max_output_bytes

    def extract(self, document_path: Path, document_format: str) -> str:
        """Extract one document using a format-specific lazy parser."""

        normalized_format = document_format.strip().lower().lstrip(".")
        handlers: dict[str, Callable[[Path], str]] = {
            "pdf": self._extract_pdf,
            "docx": self._extract_docx,
            "xlsx": self._extract_excel,
            "xls": self._extract_excel,
            "rtf": self._extract_rtf,
            "odt": self._extract_odt,
            "pptx": self._extract_pptx,
            "epub": self._extract_epub,
        }
        if sys.platform == "win32":
            handlers.update({"doc": self._extract_doc, "ppt": self._extract_ppt})
        handler = handlers.get(normalized_format)
        if handler is None:
            raise ValueError(f"Unsupported document format '{normalized_format}'.")
        try:
            text = handler(document_path)
        except Exception as error:
            if isinstance(error, (RuntimeError, ValueError)):
                raise
            raise RuntimeError(
                f"{normalized_format.upper()} document extraction failed: {error}"
            ) from error
        return self._validate_output(text)

    def _extract_pdf(self, document_path: Path) -> str:
        """Extract page text from one PDF file."""

        from pypdf import PdfReader

        reader = PdfReader(str(document_path))
        return self._join_bounded((page.extract_text() or "" for page in reader.pages), "\n")

    def _extract_docx(self, document_path: Path) -> str:
        """Extract paragraphs and tables from one DOCX file."""

        from docx import Document

        document = Document(str(document_path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return self._join_bounded(parts, "\n")

    def _extract_excel(self, document_path: Path) -> str:
        """Extract visible XLSX/XLSM sheets, falling back to legacy XLS parsing."""

        try:
            return self._extract_openpyxl(document_path)
        except Exception as xlsx_error:
            try:
                return self._extract_xlrd(document_path)
            except Exception as xls_error:
                raise RuntimeError(
                    "Excel document extraction failed. "
                    f"XLSX parser: {xlsx_error}; XLS parser: {xls_error}"
                ) from xls_error

    def _extract_openpyxl(self, document_path: Path) -> str:
        """Extract workbook values with OpenPyXL in read-only mode."""

        from openpyxl import load_workbook

        workbook = load_workbook(filename=str(document_path), read_only=True, data_only=True)
        sheets: list[str] = []
        try:
            for sheet in workbook:
                if sheet.sheet_state == "hidden":
                    continue
                rows = [f"=== Sheet: {sheet.title} ==="]
                row_count = 0
                for row in sheet.iter_rows(values_only=True):
                    if not any(cell is not None and cell != "" for cell in row):
                        continue
                    rows.append("\t".join(str(cell) if cell is not None else "" for cell in row))
                    row_count += 1
                if row_count > 0:
                    sheets.append(self._join_bounded(rows, "\n"))
            return self._join_bounded(sheets, "\n\n")
        finally:
            workbook.close()

    def _extract_xlrd(self, document_path: Path) -> str:
        """Extract legacy XLS cell values with xlrd."""

        import xlrd

        workbook = xlrd.open_workbook(filename=str(document_path), formatting_info=False)
        sheets: list[str] = []
        for sheet in workbook.sheets():
            if sheet.nrows == 0:
                continue
            rows = [f"=== Sheet: {sheet.name} ==="]
            for row_index in range(sheet.nrows):
                row_text = "\t".join(str(cell) for cell in sheet.row_values(row_index))
                if row_text.strip():
                    rows.append(row_text)
            sheets.append(self._join_bounded(rows, "\n"))
        return self._join_bounded(sheets, "\n\n")

    def _extract_rtf(self, document_path: Path) -> str:
        """Convert one UTF-8-compatible RTF document to plain text."""

        from striprtf.striprtf import rtf_to_text

        source = document_path.read_bytes().decode("utf-8", errors="replace")
        return rtf_to_text(source)

    def _extract_odt(self, document_path: Path) -> str:
        """Extract paragraphs and table rows from one ODT document."""

        from odf import table, text
        from odf.opendocument import load
        from odf.teletype import extractText

        document = load(str(document_path))
        parts = [extractText(paragraph) for paragraph in document.getElementsByType(text.P)]
        for document_table in document.getElementsByType(table.Table):
            for row in document_table.getElementsByType(table.TableRow):
                cells = [extractText(cell) for cell in row.getElementsByType(table.TableCell)]
                parts.append("\t".join(cells))
        return self._join_bounded(parts, "\n")

    def _extract_pptx(self, document_path: Path) -> str:
        """Extract shape text and tables from one PPTX presentation."""

        from pptx import Presentation

        presentation = Presentation(str(document_path))
        parts: list[str] = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                shape_text = getattr(shape, "text", "")
                if isinstance(shape_text, str) and shape_text.strip():
                    parts.append(shape_text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append("\t".join(cell.text_frame.text.strip() for cell in row.cells))
        return self._join_bounded(parts, "\n")

    def _extract_epub(self, document_path: Path) -> str:
        """Extract spine-ordered EPUB chapters as the legacy JSON string contract."""

        with zipfile.ZipFile(document_path, "r") as epub_zip:
            self._validate_epub_archive(epub_zip)
            container_root = self._read_xml(epub_zip, "META-INF/container.xml")
            root_file = container_root.find(".//{*}rootfile")
            if root_file is None or not root_file.get("full-path"):
                raise ValueError("EPUB root package path was not found.")
            opf_path = self._normalize_epub_member(root_file.get("full-path", ""))
            opf_root = self._read_xml(epub_zip, opf_path)
            namespace = {"opf": "http://www.idpf.org/2007/opf"}
            spine = opf_root.find(".//opf:spine", namespace)
            if spine is None:
                raise ValueError("EPUB spine element was not found.")
            item_references = [
                item.get("idref") for item in spine.findall("opf:itemref", namespace)
            ]
            manifest = {
                item.get("id"): posixpath.normpath(item.get("href", ""))
                for item in opf_root.findall(".//opf:item", namespace)
                if item.get("id") and item.get("href")
            }
            opf_directory = posixpath.dirname(opf_path)
            casefold_names = {name.replace("\\", "/").lower(): name for name in epub_zip.namelist()}
            chapters: list[str] = []
            processed: set[str] = set()
            for item_id in item_references:
                relative_path = manifest.get(item_id)
                if not relative_path:
                    continue
                joined_path = posixpath.join(opf_directory, relative_path) if opf_directory else relative_path
                normalized_path = self._normalize_epub_member(joined_path)
                actual_path = casefold_names.get(normalized_path.lower())
                if actual_path is None or actual_path in processed:
                    continue
                title, body = self._parse_epub_chapter(epub_zip.read(actual_path))
                chapter = f"{title}\n\n{body}" if title else body
                if chapter.strip():
                    chapters.append(chapter)
                processed.add(actual_path)
            result = json.dumps({"chapters": chapters}, ensure_ascii=False)
            return self._validate_output(result)

    def _extract_doc(self, document_path: Path) -> str:
        """Extract one legacy DOC file through a hidden Windows Word instance."""

        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        word = None
        document = None
        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            document = word.Documents.Open(str(document_path), ReadOnly=True)
            return str(document.Range().Text).strip()
        finally:
            if document is not None:
                document.Close(False)
            if word is not None:
                word.Quit()
            pythoncom.CoUninitialize()

    def _extract_ppt(self, document_path: Path) -> str:
        """Extract one legacy PPT file through a hidden Windows PowerPoint instance."""

        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        powerpoint = None
        presentation = None
        try:
            powerpoint = win32com.client.Dispatch("PowerPoint.Application")
            presentation = powerpoint.Presentations.Open(
                str(document_path),
                ReadOnly=True,
                Untitled=False,
                WithWindow=False,
            )
            parts: list[str] = []
            for slide in presentation.Slides:
                for shape in slide.Shapes:
                    if shape.HasTextFrame:
                        parts.append(str(shape.TextFrame.TextRange.Text).strip())
            return self._join_bounded((part for part in parts if part), "\n")
        finally:
            if presentation is not None:
                presentation.Close()
            if powerpoint is not None:
                powerpoint.Quit()
            pythoncom.CoUninitialize()

    def _validate_epub_archive(self, archive: zipfile.ZipFile) -> None:
        """Reject EPUB archives with unsafe names or excessive expansion."""

        entries = archive.infolist()
        if len(entries) > DEFAULT_MAX_EPUB_ENTRIES:
            raise ValueError("EPUB archive contains too many entries.")
        total_size = 0
        for entry in entries:
            self._normalize_epub_member(entry.filename)
            total_size += entry.file_size
            if total_size > DEFAULT_MAX_EPUB_UNCOMPRESSED_BYTES:
                raise ValueError("EPUB archive exceeds the uncompressed size limit.")

    def _normalize_epub_member(self, value: str) -> str:
        """Normalize one EPUB member and reject absolute or parent traversal paths."""

        normalized = posixpath.normpath(value.replace("\\", "/"))
        path = PurePosixPath(normalized)
        if not normalized or normalized == "." or path.is_absolute() or ".." in path.parts:
            raise ValueError("EPUB archive contains an unsafe member path.")
        return normalized

    def _read_xml(self, archive: zipfile.ZipFile, member: str):
        """Read one validated EPUB XML member into an ElementTree element."""

        import xml.etree.ElementTree as element_tree

        normalized = self._normalize_epub_member(member)
        return element_tree.fromstring(archive.read(normalized))

    def _parse_epub_chapter(self, html_data: bytes) -> tuple[str, str]:
        """Extract a chapter title and body from XHTML or malformed HTML bytes."""

        import xml.etree.ElementTree as element_tree

        try:
            root = element_tree.fromstring(html_data)
            namespace = {"xhtml": "http://www.w3.org/1999/xhtml"}
            title = ""
            for level in range(1, 7):
                title_element = root.find(f".//xhtml:h{level}", namespace)
                if title_element is not None:
                    title = "".join(title_element.itertext()).strip()
                    if title:
                        break
            body_element = root.find(".//xhtml:body", namespace)
            body = "".join(body_element.itertext()).strip() if body_element is not None else ""
            if title and body.startswith(title):
                body = body[len(title):].strip()
            return title, body
        except element_tree.ParseError:
            html_text = html_data.decode("utf-8", errors="replace")
            title_match = re.search(
                r"<h[1-6][^>]*>(.*?)</h[1-6]>",
                html_text,
                re.IGNORECASE | re.DOTALL,
            )
            title = title_match.group(1).strip() if title_match else ""
            body_match = re.search(
                r"<body[^>]*>(.*?)</body>",
                html_text,
                re.IGNORECASE | re.DOTALL,
            )
            body_source = body_match.group(1) if body_match else html_text
            body = re.sub(r"<[^>]+>", "", body_source).strip()
            return title, body

    def _join_bounded(self, parts, separator: str) -> str:
        """Join text parts and fail before the UTF-8 output grows beyond its limit."""

        accepted: list[str] = []
        current_bytes = 0
        separator_bytes = len(separator.encode("utf-8"))
        for part in parts:
            text = str(part)
            added_bytes = len(text.encode("utf-8"))
            if accepted:
                added_bytes += separator_bytes
            current_bytes += added_bytes
            if current_bytes > self._max_output_bytes:
                raise ValueError("Extracted document text exceeds the output size limit.")
            accepted.append(text)
        return separator.join(accepted)

    def _validate_output(self, text: str) -> str:
        """Validate extracted text type and final UTF-8 byte count."""

        if not isinstance(text, str):
            raise TypeError("Document parsers must return text.")
        if len(text.encode("utf-8")) > self._max_output_bytes:
            raise ValueError("Extracted document text exceeds the output size limit.")
        return text
