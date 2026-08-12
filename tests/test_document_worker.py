# -*- coding: utf-8 -*-
"""Tests for Document Worker artifacts, parsers, and compatibility client."""

from __future__ import annotations

from collections.abc import Mapping
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from typing import Any
import zipfile

from py.document_worker_client import DocumentWorkerClient
from py.workers.document_engine import DocumentEngine
from py.workers.document_worker import DocumentWorkerHandlers


class FakeDocumentEngine:
    """Small engine double returning document bytes as text."""

    def extract(self, document_path: Path, document_format: str) -> str:
        """Return deterministic text containing format and artifact content."""

        return f"{document_format}:{document_path.read_text(encoding='utf-8')}"


class DocumentWorkerHandlerTests(unittest.IsolatedAsyncioTestCase):
    """Validate Worker artifact and parser dependency boundaries."""

    async def test_extract_writes_text_to_result_artifact(self) -> None:
        """Keep document bytes and extracted text outside the JSON protocol."""

        with tempfile.TemporaryDirectory(prefix="openxnet-document-worker-") as directory:
            root = Path(directory)
            input_path = root / "input.rtf"
            result_path = root / "result.txt"
            input_path.write_text("document", encoding="utf-8")
            result_path.write_bytes(b"")
            handlers = DocumentWorkerHandlers(
                root,
                engine=FakeDocumentEngine(),  # type: ignore[arg-type]
            )

            response = await handlers.extract(
                {
                    "artifactPath": str(input_path),
                    "resultArtifactPath": str(result_path),
                    "format": "rtf",
                }
            )

            self.assertEqual(result_path.read_text(encoding="utf-8"), "rtf:document")
            self.assertEqual(response["resultBytes"], len("rtf:document"))

    async def test_rejects_artifacts_outside_exchange_root(self) -> None:
        """Prevent Document Worker from reading arbitrary local files."""

        with tempfile.TemporaryDirectory(prefix="openxnet-document-worker-") as directory:
            root = Path(directory)
            exchange_root = root / "exchange"
            exchange_root.mkdir()
            outside_path = root / "outside.pdf"
            result_path = exchange_root / "result.txt"
            outside_path.write_bytes(b"document")
            result_path.write_bytes(b"")
            handlers = DocumentWorkerHandlers(exchange_root, engine=FakeDocumentEngine())  # type: ignore[arg-type]

            with self.assertRaisesRegex(ValueError, "outside"):
                await handlers.extract(
                    {
                        "artifactPath": str(outside_path),
                        "resultArtifactPath": str(result_path),
                        "format": "pdf",
                    }
                )

    async def test_status_does_not_import_document_parsers(self) -> None:
        """Keep all optional parser libraries unloaded during readiness checks."""

        dependency_names = ("pypdf", "docx", "openpyxl", "xlrd", "striprtf", "odf", "pptx")
        before = {name for name in dependency_names if name in sys.modules}
        with tempfile.TemporaryDirectory(prefix="openxnet-document-status-") as directory:
            status = DocumentWorkerHandlers(Path(directory)).status({})
        after = {name for name in dependency_names if name in sys.modules}

        self.assertEqual(after, before)
        self.assertIn("pdf", status["supportedFormats"])
        self.assertIn("epub", status["availableFormats"])


class RecordingDocumentWorkerClient(DocumentWorkerClient):
    """Client double recording live artifacts instead of opening HTTP."""

    def __init__(self, exchange_root: Path) -> None:
        """Create a configured client bound to one temporary directory."""

        super().__init__("http://127.0.0.1:1", "token", exchange_root)
        self.input_path: Path | None = None
        self.result_path: Path | None = None

    def _request_sync(self, method: str, payload: Mapping[str, Any]) -> Mapping[str, Any]:
        """Read the input artifact and write deterministic extracted text."""

        if method != "documents.extract":
            raise AssertionError(f"Unexpected method: {method}")
        self.input_path = Path(str(payload["artifactPath"]))
        self.result_path = Path(str(payload["resultArtifactPath"]))
        if not self.input_path.is_file() or not self.result_path.is_file():
            raise AssertionError("Document artifacts were not live during RPC.")
        output = f"worker:{self.input_path.read_bytes().decode('utf-8')}"
        self.result_path.write_text(output, encoding="utf-8")
        return {
            "resultArtifactPath": str(self.result_path),
            "resultBytes": len(output.encode("utf-8")),
        }


class DocumentWorkerClientTests(unittest.IsolatedAsyncioTestCase):
    """Validate compatibility-client artifact lifetime and response checks."""

    async def test_extract_removes_both_artifacts(self) -> None:
        """Delete input and result files after a successful extraction."""

        with tempfile.TemporaryDirectory(prefix="openxnet-document-client-") as directory:
            client = RecordingDocumentWorkerClient(Path(directory))
            text = await client.extract(b"document", "rtf")
            input_path = client.input_path
            result_path = client.result_path

        self.assertEqual(text, "worker:document")
        self.assertIsNotNone(input_path)
        self.assertIsNotNone(result_path)
        self.assertFalse(input_path.exists())
        self.assertFalse(result_path.exists())

    async def test_legacy_office_handler_delegates_to_document_worker(self) -> None:
        """Keep the old load-files entrypoint without running parsers in the backend."""

        from unittest.mock import AsyncMock, patch

        with patch(
            "py.document_worker_client.extract_office_document",
            new_callable=AsyncMock,
            return_value="delegated document",
        ) as extract:
            from py.load_files import handle_office_document

            result = await handle_office_document(b"document", "pdf")

        self.assertEqual(result, "delegated document")
        extract.assert_awaited_once_with(b"document", "pdf")


class DocumentEngineTests(unittest.TestCase):
    """Exercise portable parsers with small deterministic documents."""

    def test_extracts_rtf(self) -> None:
        """Preserve plain text extracted from one RTF document."""

        if importlib.util.find_spec("striprtf") is None:
            self.skipTest("striprtf is not installed")
        with tempfile.TemporaryDirectory(prefix="openxnet-document-engine-") as directory:
            root = Path(directory)
            rtf_path = root / "sample.rtf"
            rtf_path.write_text(r"{\rtf1\ansi Hello document}", encoding="utf-8")
            rtf_text = DocumentEngine().extract(rtf_path, "rtf")

        self.assertIn("Hello document", rtf_text)

    def test_extracts_epub(self) -> None:
        """Preserve spine order and the legacy EPUB JSON string contract."""

        with tempfile.TemporaryDirectory(prefix="openxnet-document-engine-") as directory:
            epub_path = Path(directory) / "sample.epub"
            self._write_epub(epub_path)

            epub_value = json.loads(DocumentEngine().extract(epub_path, "epub"))

        self.assertEqual(epub_value["chapters"], ["Chapter One\n\nBody text"])

    def test_extracts_generated_open_xml_documents(self) -> None:
        """Extract generated DOCX, XLSX, and PPTX fixtures when dependencies exist."""

        dependencies = ("docx", "openpyxl", "pptx", "odf")
        if any(importlib.util.find_spec(name) is None for name in dependencies):
            self.skipTest("Open XML document parsers are not installed")
        from docx import Document
        from openpyxl import Workbook
        from odf.opendocument import OpenDocumentText
        from odf.text import P
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory(prefix="openxnet-document-openxml-") as directory:
            root = Path(directory)
            docx_path = root / "sample.docx"
            xlsx_path = root / "sample.xlsx"
            pptx_path = root / "sample.pptx"
            odt_path = root / "sample.odt"

            document = Document()
            document.add_paragraph("DOCX paragraph")
            document.save(docx_path)

            workbook = Workbook()
            workbook.active["A1"] = "XLSX cell"
            workbook.save(xlsx_path)
            workbook.close()

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            text_box.text = "PPTX shape"
            presentation.save(pptx_path)

            odt_document = OpenDocumentText()
            odt_document.text.addElement(P(text="ODT paragraph"))
            odt_document.save(str(odt_path), addsuffix=False)

            engine = DocumentEngine()
            results = {
                "docx": engine.extract(docx_path, "docx"),
                "xlsx": engine.extract(xlsx_path, "xlsx"),
                "pptx": engine.extract(pptx_path, "pptx"),
                "odt": engine.extract(odt_path, "odt"),
            }

        self.assertIn("DOCX paragraph", results["docx"])
        self.assertIn("XLSX cell", results["xlsx"])
        self.assertIn("PPTX shape", results["pptx"])
        self.assertIn("ODT paragraph", results["odt"])

    def test_extracts_blank_pdf(self) -> None:
        """Load a generated PDF through the maintained pypdf boundary."""

        if importlib.util.find_spec("pypdf") is None:
            self.skipTest("pypdf is not installed")
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory(prefix="openxnet-document-pdf-") as directory:
            pdf_path = Path(directory) / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=72, height=72)
            with pdf_path.open("wb") as output:
                writer.write(output)
            result = DocumentEngine().extract(pdf_path, "pdf")

        self.assertEqual(result, "")

    def _write_epub(self, path: Path) -> None:
        """Write one minimal EPUB archive with a single XHTML spine item."""

        container = """<?xml version="1.0"?>
<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles>
</container>"""
        package = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf">
  <manifest><item id="chapter" href="chapter.xhtml"/></manifest>
  <spine><itemref idref="chapter"/></spine>
</package>"""
        chapter = """<html xmlns="http://www.w3.org/1999/xhtml"><body>
<h1>Chapter One</h1><p>Body text</p></body></html>"""
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("META-INF/container.xml", container)
            archive.writestr("OEBPS/content.opf", package)
            archive.writestr("OEBPS/chapter.xhtml", chapter)


if __name__ == "__main__":
    unittest.main()
