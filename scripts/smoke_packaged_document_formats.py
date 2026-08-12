# -*- coding: utf-8 -*-
"""Exercise every portable format through a packaged Document Worker executable."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import subprocess
import tempfile
from typing import Any
from uuid import uuid4
import zipfile


def parse_arguments() -> argparse.Namespace:
    """Parse the packaged Document Worker executable path."""

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("worker_executable", type=Path)
    return parser.parse_args()


def write_fixtures(root: Path) -> dict[str, tuple[Path, str]]:
    """Generate small deterministic documents and their expected extracted text."""

    from docx import Document
    from odf.opendocument import OpenDocumentText
    from odf.text import P
    from openpyxl import Workbook
    from pypdf import PdfWriter
    from pptx import Presentation
    from pptx.util import Inches

    fixtures: dict[str, tuple[Path, str]] = {}

    pdf_path = root / "sample.pdf"
    pdf_writer = PdfWriter()
    pdf_writer.add_blank_page(width=72, height=72)
    with pdf_path.open("wb") as output:
        pdf_writer.write(output)
    fixtures["pdf"] = (pdf_path, "")

    docx_path = root / "sample.docx"
    document = Document()
    document.add_paragraph("Packaged DOCX text")
    document.save(docx_path)
    fixtures["docx"] = (docx_path, "Packaged DOCX text")

    xlsx_path = root / "sample.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Packaged XLSX text"
    workbook.save(xlsx_path)
    workbook.close()
    fixtures["xlsx"] = (xlsx_path, "Packaged XLSX text")

    odt_path = root / "sample.odt"
    odt_document = OpenDocumentText()
    odt_document.text.addElement(P(text="Packaged ODT text"))
    odt_document.save(str(odt_path), addsuffix=False)
    fixtures["odt"] = (odt_path, "Packaged ODT text")

    pptx_path = root / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    text_box.text = "Packaged PPTX text"
    presentation.save(pptx_path)
    fixtures["pptx"] = (pptx_path, "Packaged PPTX text")

    rtf_path = root / "sample.rtf"
    rtf_path.write_text(r"{\rtf1\ansi Packaged RTF text}", encoding="utf-8")
    fixtures["rtf"] = (rtf_path, "Packaged RTF text")

    epub_path = root / "sample.epub"
    write_epub(epub_path)
    fixtures["epub"] = (epub_path, "Packaged EPUB text")
    return fixtures


def write_epub(path: Path) -> None:
    """Write one minimal EPUB archive containing a single chapter."""

    container = """<?xml version="1.0"?>
<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles>
</container>"""
    package = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf">
  <manifest><item id="chapter" href="chapter.xhtml"/></manifest>
  <spine><itemref idref="chapter"/></spine>
</package>"""
    chapter = """<html xmlns="http://www.w3.org/1999/xhtml"><body>
<h1>Chapter</h1><p>Packaged EPUB text</p></body></html>"""
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", package)
        archive.writestr("OEBPS/chapter.xhtml", chapter)


def request_worker(
    process: subprocess.Popen[str],
    method: str,
    payload: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Send one protocol request and return its successful payload."""

    if process.stdin is None or process.stdout is None:
        raise RuntimeError("Packaged Document Worker pipes are unavailable.")
    message_id = str(uuid4())
    envelope = {
        "protocolVersion": "1.0",
        "messageId": message_id,
        "traceId": message_id,
        "kind": "request",
        "capability": "documents",
        "method": method,
        "payload": payload or {},
    }
    process.stdin.write(json.dumps(envelope, ensure_ascii=False) + "\n")
    process.stdin.flush()
    response_line = process.stdout.readline()
    if not response_line:
        diagnostic = process.stderr.read() if process.stderr is not None else ""
        raise RuntimeError(f"Packaged Document Worker exited without a response: {diagnostic}")
    response = json.loads(response_line)
    if response.get("messageId") != message_id or response.get("kind") != "response":
        raise RuntimeError(f"Packaged Document Worker request failed: {response}")
    result = response.get("payload")
    if not isinstance(result, dict):
        raise RuntimeError("Packaged Document Worker response payload is invalid.")
    return result


def main() -> int:
    """Generate fixtures, run packaged extraction, and print a JSON summary."""

    arguments = parse_arguments()
    executable = arguments.worker_executable.resolve()
    if not executable.is_file():
        raise FileNotFoundError(f"Document Worker executable was not found: {executable}")

    with tempfile.TemporaryDirectory(prefix="openxnet-packaged-document-formats-") as directory:
        exchange_root = Path(directory)
        fixtures = write_fixtures(exchange_root)
        process = subprocess.Popen(
            [str(executable), "--exchange-root", str(exchange_root)],
            cwd=executable.parent,
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
        )
        results: dict[str, dict[str, Any]] = {}
        try:
            request_worker(process, "system.ping")
            for document_format, (input_path, expected_text) in fixtures.items():
                result_path = exchange_root / f"result-{document_format}.txt"
                result_path.write_bytes(b"")
                response = request_worker(
                    process,
                    "documents.extract",
                    {
                        "artifactPath": str(input_path),
                        "resultArtifactPath": str(result_path),
                        "format": document_format,
                    },
                )
                text = result_path.read_text(encoding="utf-8")
                if expected_text not in text:
                    raise RuntimeError(
                        f"Packaged {document_format} result did not contain expected text: {text!r}"
                    )
                results[document_format] = {
                    "resultBytes": response.get("resultBytes"),
                    "characters": response.get("characters"),
                }
            request_worker(process, "system.shutdown")
            process.wait(timeout=10)
        finally:
            if process.poll() is None:
                process.kill()
                process.wait(timeout=5)
        print(json.dumps({"formats": results, "count": len(results)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
